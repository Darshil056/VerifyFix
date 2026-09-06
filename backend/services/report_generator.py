import io
import datetime
from typing import Dict, Any
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def generate_docx_report(state: Dict[str, Any]) -> io.BytesIO:
    """
    Generates a professional .docx vulnerability report using python-docx.
    """
    doc = Document()

    # Document properties
    scan_id = state.get("scan_id", "Unknown")
    repo_owner = state.get("repo_owner", "Unknown")
    repo_name = state.get("repo_name", "Unknown")
    scan_timestamp = state.get("completed_at", datetime.datetime.now().isoformat())

    # --- Title & Cover ---
    title = doc.add_heading("VerifyFix AI Vulnerability Validation Report", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_paragraph()
    
    meta_p = doc.add_paragraph()
    meta_p.add_run("Target Repository: ").bold = True
    meta_p.add_run(f"{repo_owner}/{repo_name}\n")
    meta_p.add_run("Scan ID: ").bold = True
    meta_p.add_run(f"{scan_id}\n")
    meta_p.add_run("Date: ").bold = True
    meta_p.add_run(f"{scan_timestamp}\n")
    meta_p.add_run("Authors: ").bold = True
    meta_p.add_run("Darshil Mendapara & Dhruv Parsana")

    doc.add_page_break()

    # --- Executive Summary ---
    doc.add_heading("Executive Summary", level=1)
    
    summary_data = state.get("summary", {})
    confirmed_count = summary_data.get("confirmed_vulnerabilities", 0)
    secure_count = summary_data.get("secure_findings", 0)
    pruned_count = state.get("pruned_vulns_count", 0) - len(state.get("candidate_vulns", [])) if "pruned_vulns_count" in state else 0
    # ensure it's positive
    pruned_count = abs(pruned_count)
    
    total_candidates = len(state.get("candidate_vulns", []))
    
    exec_p = doc.add_paragraph(
        f"VerifyFix performed an automated security analysis on {repo_owner}/{repo_name}. "
        f"The Discovery Agent identified {total_candidates} candidate vulnerabilities. "
        f"The Critic Agent pruned {pruned_count} false positives. "
        f"The Sandbox Engine executed test fixtures, confirming {confirmed_count} vulnerabilities "
        f"and proving {secure_count} to be secure."
    )
    
    # Summary Table
    table = doc.add_table(rows=1, cols=2)
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Metric'
    hdr_cells[1].text = 'Value'
    
    metrics = [
        ("Total Candidates", str(total_candidates)),
        ("Pruned False Positives", str(pruned_count)),
        ("Confirmed Vulnerabilities", str(confirmed_count)),
        ("Secure / Mitigated", str(secure_count)),
        ("Self-Healing Retries", str(summary_data.get("retry_count", 0)))
    ]
    
    for metric, value in metrics:
        row_cells = table.add_row().cells
        row_cells[0].text = metric
        row_cells[1].text = value
        
    doc.add_paragraph()

    # --- Detailed Findings ---
    doc.add_heading("Detailed Findings & Remediation", level=1)
    
    remediations = state.get("final_remediations", [])
    
    if not remediations:
        doc.add_paragraph("No confirmed vulnerabilities requiring remediation were found.")
    
    for i, rem in enumerate(remediations, 1):
        cwe_id = rem.get("cwe_id", "Unknown CWE")
        v_name = rem.get("vulnerability_name", "Unknown Vulnerability")
        severity = rem.get("severity", "Unknown")
        file_path = rem.get("file_path", "Unknown File")
        line_range = rem.get("line_range", "Unknown Lines")
        
        doc.add_heading(f"Finding #{i}: [{cwe_id}] {v_name}", level=2)
        
        info_p = doc.add_paragraph()
        info_p.add_run("Severity: ").bold = True
        info_p.add_run(f"{severity}\n")
        info_p.add_run("Location: ").bold = True
        info_p.add_run(f"{file_path} (Lines: {line_range})\n")
        
        # Root Cause
        doc.add_heading("Root Cause Analysis", level=3)
        doc.add_paragraph(rem.get("root_cause", "No root cause provided."))
        
        # Execution Proof
        doc.add_heading("Execution Proof (Sandbox Output)", level=3)
        proof_p = doc.add_paragraph(rem.get("execution_proof", "No proof provided."))
        # Basic monospace styling
        for run in proof_p.runs:
            run.font.name = 'Courier New'
            run.font.size = Pt(9)
            
        # Patch Diff
        doc.add_heading("Recommended Patch (Git Diff)", level=3)
        diff_p = doc.add_paragraph(rem.get("patch_diff", "No patch provided."))
        for run in diff_p.runs:
            run.font.name = 'Courier New'
            run.font.size = Pt(9)
            
        # Defense in Depth
        doc.add_heading("Defense in Depth", level=3)
        for advice in rem.get("defense_in_depth", []):
            doc.add_paragraph(advice, style='List Bullet')
            
        if i < len(remediations):
            doc.add_page_break()

    # Save to BytesIO
    f = io.BytesIO()
    doc.save(f)
    f.seek(0)
    return f
